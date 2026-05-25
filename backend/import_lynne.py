"""Import LYNNE product catalog from Produktaufstellung_KWxx.xlsx.

Reads sheet "Verkäufe", upserts into `lynne_products`. Idempotent: re-running
with a newer KW xlsx replaces stock/sales values for the same id.

Usage:
    .venv/bin/python -m backend.import_lynne \\
        --source Produktaufstellung_KW14_-_Aktuel.xlsx

Column layout (header row index 4, 0-indexed; data starts row index 5):
    col 8  OHNE LST (actually = ASIN, e.g. "B092R3GC2L")
    col 10 PRIME/Eigen Versand (channel: PRIME / EV / EV-PRIME / OTHER)
    col 11 SKU
    col 12 Interne Artikel Nummer (EAN)
    col 13 Artikel Beschreibung
    col 14 Marke
    col 16 Wöchentliche Verkäufe (Vorwoche)
    col 19 Wie auf die Palette
    col 20 Lagerbestand Graz 6 Monate
"""

from __future__ import annotations

import argparse
import asyncio
import re
import sys
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert

from backend.database import AsyncSessionLocal
from backend.orm import LynneProduct


ASIN_RE = re.compile(r"^B0[A-Z0-9]{8}$", re.IGNORECASE)

COL_ASIN = 8
COL_CHANNEL = 10
COL_SKU = 11
COL_EAN = 12
COL_DESCRIPTION = 13
COL_BRAND = 14
COL_WEEKLY_SALES = 16
COL_PER_PALLET = 19
COL_GRAZ_STOCK = 20

DATA_START_ROW = 5  # 0-indexed; header is row 4


def _str(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _ean(v) -> str | None:
    """Accept only pure digits 8-14 chars (real EAN/UPC). Mining the
    xlsx shows the EAN column often contains free-form notes
    ("wird von 9120107... produziert", "Nicht vorrätig", etc.) —
    those drop to None."""
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    if s.isdigit() and 8 <= len(s) <= 14:
        return s
    return None


def _int(v) -> int:
    """Parse cell to int, tolerating floats ('560.0'), commas, and blanks."""
    if v is None or v == "":
        return 0
    try:
        s = str(v).replace(",", ".").strip()
        f = float(s)
        return int(round(f))
    except (ValueError, TypeError):
        return 0


def _normalize_channel(raw: str) -> str:
    s = raw.upper().strip()
    if s in {"PRIME", "EV", "EV-PRIME", "OTHER"}:
        return s
    # Accept variants like "PRIME ", "ev", "EV PRIME"
    if "PRIME" in s and "EV" in s:
        return "EV-PRIME"
    if s == "PRIME":
        return "PRIME"
    if s == "EV":
        return "EV"
    return "OTHER"


def parse_xlsx(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    if "Verkäufe" not in wb.sheetnames:
        raise SystemExit(
            f"Sheet 'Verkäufe' not found in {path.name}. "
            f"Available sheets: {wb.sheetnames}"
        )
    ws = wb["Verkäufe"]
    rows = []
    seen_ids: set[str] = set()

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < DATA_START_ROW:
            continue
        if not row or len(row) <= COL_GRAZ_STOCK:
            continue

        asin = _str(row[COL_ASIN])
        if not ASIN_RE.match(asin):
            continue
        sku = _str(row[COL_SKU])
        if not sku:
            continue

        pid = f"{asin}__{sku}"
        if pid in seen_ids:
            continue
        seen_ids.add(pid)

        rows.append({
            "id": pid,
            "asin": asin,
            "channel": _normalize_channel(_str(row[COL_CHANNEL])),
            "sku": sku,
            "ean": _ean(row[COL_EAN]),
            "description": _str(row[COL_DESCRIPTION]),
            "brand": _str(row[COL_BRAND]),
            "weekly_sales": _int(row[COL_WEEKLY_SALES]),
            "per_pallet": _int(row[COL_PER_PALLET]),
            "graz_stock": _int(row[COL_GRAZ_STOCK]),
            "source": f"xlsx:{path.name}",
        })

    wb.close()
    return rows


async def upsert(rows: list[dict]) -> tuple[int, int]:
    """Returns (inserted, updated)."""
    if not rows:
        return (0, 0)

    async with AsyncSessionLocal() as db:
        existing_ids = set(
            (await db.execute(select(LynneProduct.id))).scalars().all()
        )
        inserted = sum(1 for r in rows if r["id"] not in existing_ids)
        updated = len(rows) - inserted

        stmt = pg_insert(LynneProduct).values(rows)
        stmt = stmt.on_conflict_do_update(
            index_elements=["id"],
            set_={
                "asin": stmt.excluded.asin,
                "channel": stmt.excluded.channel,
                "sku": stmt.excluded.sku,
                "ean": stmt.excluded.ean,
                "description": stmt.excluded.description,
                "brand": stmt.excluded.brand,
                "weekly_sales": stmt.excluded.weekly_sales,
                "per_pallet": stmt.excluded.per_pallet,
                "graz_stock": stmt.excluded.graz_stock,
                "source": stmt.excluded.source,
            },
        )
        await db.execute(stmt)
        await db.commit()

    return (inserted, updated)


async def main_async(source: Path) -> None:
    if not source.exists():
        raise SystemExit(f"Source file not found: {source}")
    print(f"Reading {source} …")
    rows = parse_xlsx(source)
    print(f"Parsed {len(rows)} valid product rows.")
    inserted, updated = await upsert(rows)
    print(f"Inserted {inserted}, updated {updated}.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source", "-s",
        type=Path,
        required=True,
        help="Path to Produktaufstellung_KWxx_-_Aktuel.xlsx",
    )
    args = parser.parse_args()
    asyncio.run(main_async(args.source))


if __name__ == "__main__":
    main()
